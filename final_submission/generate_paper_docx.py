"""增强版论文转换器：A4版式、页码、图表题注、三线表、公式编号、匿名与元数据清理。"""
from __future__ import annotations
import re
import sys
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from PIL import Image

TEXT_WIDTH_CM = 15.4
MAX_FIG_W = 14.5
MAX_FIG_H = 12.8
TABLE_CAPTIONS = [
    '主要符号表', '数学校验结果', '本地严格配对实验结果', '官方100局演练主要指标',
    '两两Mann-Whitney检验结果', '与传统方法和旧方案对比', '成本分解', '正式测试记录表',
    '总体路线比较', '候选定位方法比较', 'G25O版本对比', '代码模块清单', '支撑材料清单',
]

def set_run(run, size=12, east='宋体', ascii_f='Times New Roman', bold=False, italic=False, color=None, name=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if name:
        run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), ascii_f if not name else name)
    rfonts.set(qn('w:hAnsi'), ascii_f if not name else name)
    rfonts.set(qn('w:eastAsia'), east if not name else name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def style_fonts(style, size, east='宋体', ascii_f='Times New Roman', bold=False, name=None):
    style.font.size = Pt(size)
    style.font.bold = bold
    if name:
        style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.insert(0, rfonts)
    rfonts.set(qn('w:ascii'), ascii_f if not name else name)
    rfonts.set(qn('w:hAnsi'), ascii_f if not name else name)
    rfonts.set(qn('w:eastAsia'), east if not name else name)

def setup_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.8)
    sec.top_margin = Cm(2.8)
    sec.bottom_margin = Cm(2.8)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)
    # 正文样式
    normal = doc.styles['Normal']
    style_fonts(normal, 12, east='宋体', ascii_f='Times New Roman')
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.22
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    try:
        pf.widow_control = True
    except Exception:
        pass
    # 标题样式
    for name, size, east in [('Heading 1', 15, '黑体'), ('Heading 2', 13, '黑体'), ('Heading 3', 12, '宋体')]:
        st = doc.styles[name]
        style_fonts(st, size, east=east, ascii_f='Times New Roman', bold=True)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(10 if name != 'Heading 3' else 6)
        st.paragraph_format.space_after = Pt(6 if name != 'Heading 3' else 3)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 公式样式
    try:
        fstyle = doc.styles['FormulaPara']
    except KeyError:
        fstyle = doc.styles.add_style('FormulaPara', WD_STYLE_TYPE.PARAGRAPH)
        fstyle.base_style = doc.styles['Normal']
        style_fonts(fstyle, 11.5, east='Cambria Math', ascii_f='Cambria Math')
        fstyle.paragraph_format.first_line_indent = Cm(0)
        fstyle.paragraph_format.space_before = Pt(4)
        fstyle.paragraph_format.space_after = Pt(4)
        fstyle.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    # 页脚页码，仅页码，居中
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in list(p.runs): r._element.getparent().remove(r._element)
    run = p.add_run()
    set_run(run, 10.5, east='宋体', ascii_f='Times New Roman')
    fld1 = OxmlElement('w:fldChar'); fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fld2 = OxmlElement('w:fldChar'); fld2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
    # 页眉留空
    hdr = sec.header
    hdr.is_linked_to_previous = False
    for hp in hdr.paragraphs:
        for r in list(hp.runs): r._element.getparent().remove(r._element)

def add_body_paragraph(doc, text, first_indent=True, align=None, size=12, bold=False, east='宋体', keep_next=False, space_after=0, line_spacing=None):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Pt(size * 2) if first_indent else Cm(0)
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_next
    pf.widow_control = True
    if line_spacing is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_spacing
    r = p.add_run(text)
    set_run(r, size, east=east, bold=bold)
    return p

def set_cell_text(cell, text, bold=False, size=10.5):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    set_run(r, size, east='宋体', bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'bottom'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '12'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    for edge in ('left', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), 'auto')
        borders.append(el)
    tblPr.append(borders)

def set_header_bottom_border(row):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:space'), '0'); bottom.set(qn('w:color'), '000000')
        tcBorders.append(bottom)
        tcPr.append(tcBorders)

def add_three_line_table(doc, rows, caption):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    set_run(r, 10.5, east='宋体', bold=True)
    ncol = max(len(r_) for r_ in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j in range(ncol):
            txt = row[j].strip() if j < len(row) else ''
            txt = re.sub(r'\s+', ' ', txt)
            set_cell_text(table.cell(i, j), txt, bold=(i == 0))
        trPr = table.rows[i]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
        if i == 0:
            tblHeader = OxmlElement('w:tblHeader'); trPr.append(tblHeader)
    set_header_bottom_border(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

FUNCTION_WORDS = {'sin', 'cos', 'tan', 'max', 'min', 'log', 'exp', 'arg', 'E'}
DESC_SUB = {'meas', 'sw', 'fail', 'succ', 'vis'}

def _formula_segments(s):
    """把线性公式拆成 (文本, 上标, 下标, 斜体) 片段。"""
    out = []
    i = 0
    n = len(s)
    def base_segments(text):
        import re as _re
        pos = 0
        for m in _re.finditer(r'[A-Za-z]+|[0-9]+(?:\.[0-9]+)?|[\u0370-\u03ff]|[^A-Za-z0-9\u0370-\u03ff]', text):
            tok = m.group(0)
            italic = tok not in FUNCTION_WORDS and all((ch.isalpha() or '\u0370' <= ch <= '\u03ff') for ch in tok)
            out.append((tok, False, False, italic))
    while i < n:
        ch = s[i]
        if ch in '^_':
            sub = (ch == '_')
            i += 1
            if i >= n: break
            if s[i] == '{':
                depth = 1; j = i + 1
                while j < n and depth:
                    if s[j] == '{': depth += 1
                    elif s[j] == '}': depth -= 1
                    j += 1
                content = s[i+1:j-1]
                i = j
            else:
                mm = re.match(r'[A-Za-z0-9\u0370-\u03ff]+', s[i:])
                if mm:
                    content = mm.group(0)
                    i += len(content)
                else:
                    content = s[i]
                    i += 1
            import re as _re
            sup_flag = not sub
            for m in _re.finditer(r'[A-Za-z]+|[\u0370-\u03ff]|[^A-Za-z0-9\u0370-\u03ff]', content):
                tok = m.group(0)
                if sub and tok in DESC_SUB:
                    italic = False
                elif sup_flag and tok == 'T':
                    italic = False
                elif len(tok) == 1 and (tok.isalpha() or '\u0370' <= tok <= '\u03ff'):
                    italic = True
                else:
                    italic = False
                out.append((tok, sup_flag, sub, italic))
        else:
            import re as _re
            m = _re.match(r'[A-Za-z]+|[0-9]+(?:\.[0-9]+)?|[\u0370-\u03ff]|[^A-Za-z0-9\u0370-\u03ff]', s[i:])
            tok = m.group(0)
            italic = tok not in FUNCTION_WORDS and all((c.isalpha() or '\u0370' <= c <= '\u03ff') for c in tok)
            out.append((tok, False, False, italic))
            i += len(tok)
    merged = []
    for seg in out:
        if merged and (merged[-1][1], merged[-1][2], merged[-1][3]) == (seg[1], seg[2], seg[3]):
            merged[-1] = (merged[-1][0] + seg[0], seg[1], seg[2], seg[3])
        else:
            merged.append(seg)
    return merged

def add_formula_paragraph(doc, formula, number):
    p = doc.add_paragraph(style='FormulaPara')
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(4); pf.space_after = Pt(4)
    pf.keep_with_next = False
    pf.tab_stops.add_tab_stop(Cm(7.7), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(15.4), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run('\t'); set_run(r1, 12)
    for text, sup, sub, italic in _formula_segments(formula):
        r = p.add_run(text)
        set_run(r, 11.5, east='宋体', ascii_f='Cambria Math', name='Cambria Math', italic=italic)
        if sup: r.font.superscript = True
        if sub: r.font.subscript = True
    r3 = p.add_run('\t(' + str(number) + ')'); set_run(r3, 11.5, east='Times New Roman', ascii_f='Times New Roman')
    return p

def is_formula_line(s: str) -> bool:
    s = s.strip()
    if not s or len(s) > 260: return False
    if s.startswith('#') or s.startswith('![') or s.startswith('```') or s.startswith('|'):
        return False
    if re.match(r'^[-*]\s+', s) or re.match(r'^\d+\.\s+', s):
        return False
    if 'http' in s or '微信公众号' in s or '`' in s or '.py' in s or 'scripts/' in s:
        return False
    cjk = sum(1 for ch in s if '\u4e00' <= ch <= '\u9fff')
    if cjk > 0: return False
    ops = set('=≤≥∈∩∪Σ√^_·→×±∓≈<>')
    if not any(ch in ops for ch in s): return False
    if re.fullmatch(r'[\s\-:|]+', s): return False
    return True

def add_picture_fit(doc, img_path: Path, caption_number: int, alt: str):
    try:
        with Image.open(img_path) as im:
            w, h = im.size
        width_cm = min(MAX_FIG_W, MAX_FIG_H * w / h)
        if width_cm < 4: width_cm = 4
    except Exception:
        width_cm = 12
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(img_path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.keep_with_next = False
    r = cap.add_run('图 %d  %s' % (caption_number, alt))
    set_run(r, 10.5, east='宋体')
    return cap

def convert(md_path: Path, out_path: Path):
    doc = Document()
    setup_document(doc)
    text = md_path.read_text(encoding='utf-8')
    lines = text.splitlines()
    i = 0
    table_no = 0; figure_no = 0; formula_no = 0
    pending_page_break = False
    in_abstract = False
    while i < len(lines):
        line = lines[i].rstrip()
        # 代码块
        if line.startswith('```'):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run('\n'.join(code))
            set_run(r, 8.5, east='宋体', ascii_f='Consolas', name='Consolas')
            continue
        # 图片
        m = re.match(r'^!\[(.*?)\]\((.*?)\)', line)
        if m:
            alt, src = m.group(1), m.group(2)
            img = (md_path.parent / src).resolve()
            if img.exists():
                figure_no += 1
                add_picture_fit(doc, img, figure_no, alt)
            i += 1
            continue
        # 表格：仅当前行以|开头且下一行为分隔行
        if line.startswith('|') and i + 1 < len(lines):
            nxt = lines[i+1].strip()
            if nxt.startswith('|') and all(set(c.strip()) <= set('-: ') for c in nxt.strip().strip('|').split('|')):
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    cells = lines[i].strip().strip('|').split('|')
                    if not all(set(c.strip()) <= set('-: ') for c in cells):
                        rows.append(cells)
                    i += 1
                table_no += 1
                cap_text = TABLE_CAPTIONS[table_no-1] if table_no-1 < len(TABLE_CAPTIONS) else ('数据表%d' % table_no)
                add_three_line_table(doc, rows, '表 %d  %s' % (table_no, cap_text))
                continue
        # 标题与正文
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(line[2:].strip())
            set_run(r, 16, east='黑体', bold=True)
        elif line.startswith('## 摘要'):
            in_abstract = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run('摘  要')
            set_run(r, 14, east='黑体', bold=True)
        elif line.startswith('## '):
            title = line[3:].strip()
            if title.startswith('一、'):
                in_abstract = False
            p = doc.add_heading(title, level=1)
            if title.startswith('一、问题重述') or title.startswith('参考文献') or title.startswith('附录'):
                p.paragraph_format.page_break_before = True
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
        elif re.match(r'^\[\d+\]\s', line):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.first_line_indent = Cm(-0.74)
            pf.left_indent = Cm(0.74)
            pf.space_after = Pt(3)
            pf.widow_control = True
            r = p.add_run(line.strip())
            set_run(r, 10.5, east='宋体')
        elif line.startswith('关键词'):
            ksize = 10.5 if in_abstract else 12
            p = add_body_paragraph(doc, line.strip(), first_indent=False, size=ksize, line_spacing=(1.16 if in_abstract else None))
            for r in p.runs:
                set_run(r, ksize, east='黑体', bold=False)
        elif re.match(r'^[-*]\s+', line):
            add_body_paragraph(doc, line[2:].strip(), first_indent=False)
        elif re.match(r'^\d+\.\s+', line):
            add_body_paragraph(doc, line[line.find(' '):].strip(), first_indent=False)
        elif line.strip() == '':
            pass
        elif is_formula_line(line):
            formula_no += 1
            add_formula_paragraph(doc, line.strip(), formula_no)
        else:
            if in_abstract:
                add_body_paragraph(doc, line.strip(), size=10.5, line_spacing=1.16)
            else:
                add_body_paragraph(doc, line.strip())
        i += 1
    # 文档元数据初步置空
    cp = doc.core_properties
    cp.author = ''
    cp.last_modified_by = ''
    cp.title = ''
    cp.subject = ''
    cp.comments = ''
    cp.category = ''
    cp.keywords = ''
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print('saved', out_path, 'tables', table_no, 'figures', figure_no, 'formulas', formula_no)

if __name__ == '__main__':
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else (Path(__file__).parent / 'B题_方案B_几何保证主动搜索_论文稿.md')
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else (Path(__file__).parent / 'B题_方案B_几何保证主动搜索_论文稿_v4.docx')
    convert(src, dst)