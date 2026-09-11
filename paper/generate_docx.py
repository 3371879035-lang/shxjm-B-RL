"""将 paper.md 转换为带基础样式的 docx。"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


def set_cn_font(doc):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_md_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncol):
            text = row[j] if j < len(row) else ""
            table.cell(i, j).text = text.strip()
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path):
    doc = Document()
    set_cn_font(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("!["):
            import re as _re
            m = _re.match(r"!\[.*?\]\((.*?)\)", line)
            if m:
                img = Path(m.group(1))
                if img.exists():
                    doc.add_picture(str(img))
                    doc.add_paragraph()
            i += 1
            continue
        if line.startswith("```"):
            # 代码块
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph("\n".join(code))
            p.style = doc.styles["Normal"]
            for r in p.runs:
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            i += 1
            continue
        if line.startswith("|") and "|" in line:
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c.strip()) <= set("-: ") for c in cells):
                    rows.append(cells)
                i += 1
            add_md_table(doc, rows)
            continue
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif re.match(r"^[-*] ", line):
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(line[line.find(" "):].strip(), style="List Number")
        elif line.strip() == "":
            doc.add_paragraph()
        else:
            para = doc.add_paragraph(line)
            if line.startswith("参赛队号：") or line.startswith("队员："):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else (Path(__file__).parent / "paper.md")
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else (Path(__file__).parent / "B题_方案B_强化学习_论文稿.docx")
    convert(src, dst)
    print(dst)
